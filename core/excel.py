from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


OP_NAMES = {
    "take":         "Взял",
    "return":       "Вернул",
    "income":       "Приход",
    "usage":        "Расход",
    "writeoff":     "Списано",
    "inventory":    "Инвентаризация",
    "transfer_out": "Передал",
    "transfer_in":  "Получил",
}


def _style_header(ws, n_cols):
    fill = PatternFill("solid", fgColor="E5E7EB")
    for col in range(1, n_cols + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = fill
    ws.freeze_panes = "A2"


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_stock_excel(rows) -> bytes:
    wb = Workbook()
    headers = ["Город", "Склад", "Артикул", "Наименование", "Короткое имя",
               "Сегмент", "Применяемость", "Локация", "Ед. изм", "Остаток"]
    ws = wb.active
    ws.title = "Остатки"
    ws.append(headers)
    for r in rows:
        ws.append([
            r["city_name"] or "", r["warehouse_name"] or "",
            r["article"] or "", r["name"] or "", r["short_name"] or "",
            r["segment"] or "", r["applicability"] or "",
            r["location_name"] or "", r["unit"] or "",
            r["stock"] or 0,
        ])
    _style_header(ws, len(headers))
    _autosize(ws, [14, 20, 16, 40, 22, 10, 16, 12, 8, 10])
    return _to_bytes(wb)


def build_backpacks_excel(rows) -> bytes:
    wb = Workbook()
    headers = ["Сотрудник", "Логин", "Склад", "Артикул", "Наименование",
               "Сегмент", "Применяемость", "Ед. изм", "На руках"]
    ws = wb.active
    ws.title = "Рюкзаки"
    ws.append(headers)
    for r in rows:
        ws.append([
            r["full_name"] or "", r["username"] or "",
            r["warehouse_name"] or "",
            r["article"] or "", r["name"] or "",
            r["segment"] or "", r["applicability"] or "",
            r["unit"] or "", r["balance"] or 0,
        ])
    _style_header(ws, len(headers))
    _autosize(ws, [24, 16, 20, 16, 40, 10, 16, 8, 10])
    return _to_bytes(wb)


def build_operations_excel(rows) -> bytes:
    wb = Workbook()
    headers = ["ID", "Дата", "Сотрудник", "Логин",
               "Операция", "Артикул", "Наименование",
               "Сегмент", "Применяемость", "Склад", "Город",
               "Ед. изм", "Кол-во", "Комментарий"]
    ws = wb.active
    ws.title = "Операции"
    ws.append(headers)
    for r in rows:
        ws.append([
            r["id"], r["timestamp"],
            r["full_name"] or "", r["username"] or "",
            OP_NAMES.get(r["operation_type"], r["operation_type"]),
            r["article"] or "", r["name"] or "",
            r["segment"] or "", r["applicability"] or "",
            r["warehouse_name"] or "", r["city_name"] or "",
            r["unit"] or "", r["quantity"] or 0,
            r["note"] or "",
        ])
    _style_header(ws, len(headers))
    _autosize(ws, [6, 20, 22, 16, 16, 16, 40, 10, 16, 20, 14, 8, 8, 30])
    return _to_bytes(wb)


def build_summary_users_excel(rows) -> bytes:
    wb = Workbook()
    headers = ["Сотрудник", "Логин", "Артикул", "Наименование",
               "Сегмент", "Применяемость", "Склад", "Город",
               "Ед. изм", "Всего взял"]
    ws = wb.active
    ws.title = "По сотрудникам"
    ws.append(headers)
    for r in rows:
        ws.append([
            r["full_name"] or "", r["username"] or "",
            r["article"] or "", r["name"] or "",
            r["segment"] or "", r["applicability"] or "",
            r["warehouse_name"] or "", r["city_name"] or "",
            r["unit"] or "", r["total"] or 0,
        ])
    _style_header(ws, len(headers))
    _autosize(ws, [24, 16, 16, 40, 10, 16, 20, 14, 8, 12])
    return _to_bytes(wb)


def build_summary_items_excel(rows) -> bytes:
    wb = Workbook()
    headers = ["Город", "Склад", "Артикул", "Наименование",
               "Сегмент", "Применяемость", "Ед. изм",
               "Приход", "Выдано", "Возврат",
               "Расход", "Списано",
               "Передано", "Получено", "Инвент.", "Остаток"]
    ws = wb.active
    ws.title = "По позициям"
    ws.append(headers)
    for r in rows:
        ws.append([
            r["city_name"] or "", r["warehouse_name"] or "",
            r["article"] or "", r["name"] or "",
            r["segment"] or "", r["applicability"] or "",
            r["unit"] or "",
            r["income"] or 0, r["taken"] or 0, r["returned"] or 0,
            r["usage"] or 0, r["writeoff"] or 0,
            r["transfer_out"] or 0, r["transfer_in"] or 0,
            r["inventory"] or 0,
            r["stock"] or 0,
        ])
    _style_header(ws, len(headers))
    _autosize(ws, [14, 20, 16, 40, 10, 16, 8,
                   10, 10, 10, 10, 10, 10, 10, 10, 10])
    return _to_bytes(wb)


def build_transfers_excel(rows) -> bytes:
    wb = Workbook()
    headers = ["ID", "Создана", "Обработана", "Статус",
               "От кого", "Кому",
               "Артикул", "Наименование",
               "Сегмент", "Применяемость", "Склад", "Город",
               "Ед. изм", "Кол-во"]
    ws = wb.active
    ws.title = "Передачи"
    ws.append(headers)
    statuses = {"pending": "Ждёт", "accepted": "Принята",
                "rejected": "Отклонена", "cancelled": "Отменена"}
    for r in rows:
        ws.append([
            r["id"], r["created_at"], r["resolved_at"] or "",
            statuses.get(r["status"], r["status"]),
            r["from_name"] or "", r["to_name"] or "",
            r["article"] or "", r["name"] or "",
            r["segment"] or "", r["applicability"] or "",
            r["warehouse_name"] or "", r["city_name"] or "",
            r["unit"] or "", r["quantity"] or 0,
        ])
    _style_header(ws, len(headers))
    _autosize(ws, [6, 18, 18, 12, 22, 22, 16, 40, 10, 16, 20, 14, 8, 8])
    return _to_bytes(wb)


def _to_bytes(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_filename(prefix="report") -> str:
    return prefix + "_" + datetime.now().strftime("%Y-%m-%d_%H-%M") + ".xlsx"
