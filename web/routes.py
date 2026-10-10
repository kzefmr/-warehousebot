from io import BytesIO
from fastapi import APIRouter, Request, Form, HTTPException, UploadFile, File
from fastapi.responses import HTMLResponse, RedirectResponse, Response, FileResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from core.db import (
    # справочники
    list_cities, get_city, create_city, update_city, delete_city,
    list_warehouses, get_warehouse, create_warehouse, update_warehouse, delete_warehouse,
    list_locations, get_location, create_location, update_location, delete_location,
    # каталог
    search_catalog, list_catalog_for_warehouse, get_catalog_item,
    create_catalog_item, update_catalog_item, delete_catalog_item,
    update_catalog_sort, list_catalog, catalog_count,
    # операции
    add_operation, add_operations_bulk, add_inventory_bulk,
    list_operations, get_operation, update_operation, delete_operation,
    clear_operations,
    # пользователи
    get_user_by_username, list_users, create_user, update_user, delete_user,
    get_user_cities, set_user_cities, get_user_by_id,
    # рюкзак / передачи / остатки
    get_user_backpack, create_transfer, get_transfer,
    list_incoming_transfers, list_outgoing_transfers,
    accept_transfer, reject_transfer, cancel_transfer,
    get_warehouse_stock, get_user_backpack_stock,
    get_recent_catalog_ids, list_favorites, toggle_favorite,
    # настройки
    get_all_settings, get_bool_setting, get_setting, set_setting,
    # отчёты
    get_stock_report, get_backpacks_report, get_operations_report,
    get_summary_users_report, get_summary_items_report, get_transfers_report,
    # шкафчики
    list_locks, get_lock, get_lock_by_arduino_id,
    create_lock, update_lock, delete_lock,
    request_lock_open, cancel_lock_open, consume_lock_open,
    log_lock_event, get_lock_log, get_all_lock_log,
    get_lock_cooldown_left,
    update_lock_seen, get_lock_seen_age,
    check_unauthorized_open,
    get_lock_timings, update_lock_timings,
    # ключи
    init_empty_slots, list_keys, get_key, upsert_key,
    update_key_calibration, delete_key_data,
    find_calibration_conflicts, get_calibration_progress,
    update_slot_state, link_key_event_to_user,
    get_key_log, get_all_key_log, count_keys_in_lock,
    request_calibration, consume_calibration, is_calibration_requested,
    # инструменты
    list_tools, get_tool, get_tools_by_holder,
    next_inventory_no, create_tool, update_tool, delete_tool, set_tool_photo,
    take_tool, return_tool, repair_tool, mark_tool_lost, decommission_tool,
    create_tool_transfer, get_tool_log, get_all_tool_log, tools_count,
)
from core.security import hash_password, verify_password
from core.excel import (
    build_stock_excel, build_backpacks_excel, build_operations_excel,
    build_summary_users_excel, build_summary_items_excel, build_transfers_excel,
    export_filename,
)
from core.auth import require_user, require_admin
from core.qr import (
    make_qr_png, parse_qr_code,
    make_material_code, make_tool_code, make_key_code,
)
from core.images import save_tool_photo, get_photo_path
from core.db import get_catalog_items_by_ids
from web.deps import templates

import os
import json

router = APIRouter()

ARDUINO_TOKEN = os.getenv("ARDUINO_TOKEN", "change-me-in-amvera")


# =============== ВСПОМОГАТЕЛЬНЫЕ ===============

def _safe_int(val, default, min_v, max_v):
    try:
        n = int(val)
        return max(min_v, min(max_v, n))
    except (ValueError, TypeError):
        return default


def _one_of(val, choices, default):
    return val if val in choices else default


def _parse_dates(request: Request):
    df = request.query_params.get("date_from") or None
    dt = request.query_params.get("date_to") or None
    return df, dt


# =============== АВТОРИЗАЦИЯ ===============

@router.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    if request.session.get("user_id"):
        return RedirectResponse("/", status_code=303)
    return templates.TemplateResponse(
        "login.html",
        {"request": request, "error": None, "user": None},
    )


@router.post("/login")
async def login_submit(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
):
    user = get_user_by_username(username.strip())
    if not user or not verify_password(password, user["password_hash"]):
        return templates.TemplateResponse(
            "login.html",
            {"request": request, "error": "Неверный логин или пароль", "user": None},
            status_code=401,
        )
    request.session["user_id"] = user["id"]
    return RedirectResponse("/", status_code=303)


@router.get("/logout")
async def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


# =============== ГЛАВНАЯ ===============

@router.get("/", response_class=HTMLResponse)
async def index(request: Request):
    user = require_user(request)
    return templates.TemplateResponse("index.html", {"request": request, "user": user})


# =============== ФОТО ИНСТРУМЕНТОВ ===============

@router.get("/uploads/tools/{filename}")
async def tool_photo_file(request: Request, filename: str):
    require_user(request)
    if "/" in filename or ".." in filename:
        raise HTTPException(400, "Некорректное имя файла")
    path = get_photo_path(filename)
    if not os.path.exists(path):
        raise HTTPException(404, "Файл не найден")
    return FileResponse(path, media_type="image/jpeg")


# =============== ВЫБОР СКЛАДА ===============

@router.get("/select-warehouse", response_class=HTMLResponse)
async def select_warehouse(request: Request):
    user = require_user(request)
    user_cities = get_user_cities(user["id"])
    if not user_cities:
        return templates.TemplateResponse(
            "select_warehouse.html",
            {
                "request": request, "user": user,
                "cities": [], "warehouses_by_city": {},
                "error": "Вам не назначен ни один город. Обратитесь к администратору.",
            },
            status_code=400,
        )
    all_wh = list_warehouses()
    wh_by_city = {}
    for c in user_cities:
        wh_by_city[c["id"]] = [w for w in all_wh if w["city_id"] == c["id"]]
    return templates.TemplateResponse(
        "select_warehouse.html",
        {
            "request": request, "user": user,
            "cities": user_cities, "warehouses_by_city": wh_by_city,
            "error": None,
        },
    )


# =============== ВЗЯТЬ (материалы) ===============

@router.get("/take", response_class=HTMLResponse)
async def take_page(request: Request, warehouse_id: int):
    user = require_user(request)
    wh = get_warehouse(warehouse_id)
    if not wh:
        raise HTTPException(404, "Склад не найден")

    items = list_catalog_for_warehouse(warehouse_id, limit=1000)

    show_stock = get_bool_setting("show_stock", True)
    stock = get_warehouse_stock(warehouse_id) if show_stock else {}

    backpack_stock = get_user_backpack_stock(user["id"])

    show_fav = get_bool_setting("show_favorites", True)
    fav_ids = list_favorites(user["id"]) if show_fav else []

    show_recent = get_bool_setting("show_recent", True)
    recent_limit = int(get_setting("recent_limit", "10") or 10)
    recent_ids = get_recent_catalog_ids(user["id"], recent_limit) if show_recent else []

    fav_set = set(fav_ids)
    recent_set = set(recent_ids) - fav_set

    def sort_key(it):
        if it["id"] in fav_set:
            return (0, 0, it["sort_order"] or 1000, it["name"])
        if it["id"] in recent_set:
            return (1, 0, it["sort_order"] or 1000, it["name"])
        return (2, it["sort_order"] or 1000, it["sort_order"] or 1000, it["name"])

    items.sort(key=sort_key)

    return templates.TemplateResponse(
        "take.html",
        {
            "request": request, "user": user,
            "warehouse": wh,
            "items": items,
            "stock": stock,
            "backpack_stock": backpack_stock,
            "favorites": fav_set,
            "recent": recent_set,
            "show_stock": show_stock,
            "show_favorites": show_fav,
            "show_recent": show_recent,
        },
    )


@router.post("/take/submit")
async def take_submit(request: Request, warehouse_id: int = Form(...)):
    user = require_user(request)
    wh = get_warehouse(warehouse_id)
    if not wh:
        raise HTTPException(404, "Склад не найден")

    form = await request.form()
    cart_raw = form.get("cart", "")
    if not cart_raw:
        return RedirectResponse("/take?warehouse_id=" + str(warehouse_id), status_code=303)

    try:
        cart = json.loads(cart_raw)
    except Exception:
        raise HTTPException(400, "Некорректные данные")

    items = []
    for k, v in cart.items():
        try:
            cid = int(k)
            qty = float(v)
            if qty > 0:
                items.append((cid, qty))
        except (ValueError, TypeError):
            continue

    if items:
        add_operations_bulk(user["id"], warehouse_id, items, op_type="take")

    return RedirectResponse("/done", status_code=303)


@router.post("/favorite/toggle")
async def favorite_toggle(request: Request, catalog_id: int = Form(...), warehouse_id: int = Form(...)):
    user = require_user(request)
    toggle_favorite(user["id"], catalog_id)
    return RedirectResponse("/take?warehouse_id=" + str(warehouse_id), status_code=303)


@router.get("/done", response_class=HTMLResponse)
async def done_page(request: Request):
    user = require_user(request)
    return templates.TemplateResponse("done.html", {"request": request, "user": user})


# =============== АДМИН: УПРАВЛЕНИЕ СКЛАДОМ ===============

@router.get("/admin/stock", response_class=HTMLResponse)
async def admin_stock(request: Request, warehouse_id: int = 0, q: str = ""):
    user = require_admin(request)
    warehouses = list_warehouses()
    if not warehouse_id and warehouses:
        warehouse_id = warehouses[0]["id"]

    wh = get_warehouse(warehouse_id) if warehouse_id else None
    if not wh:
        return templates.TemplateResponse(
            "admin_stock.html",
            {
                "request": request, "user": user,
                "warehouses": warehouses, "warehouse": None,
                "items": [], "stock": {}, "query": q,
                "all_users": list_users(),
            },
        )

    items = list_catalog_for_warehouse(warehouse_id, limit=1000)
    if q:
        ql = q.lower()
        items = [it for it in items
                 if ql in (it["name"] or "").lower()
                 or ql in (it["short_name"] or "").lower()
                 or ql in (it["article"] or "").lower()
                 or ql in (it["segment"] or "").lower()
                 or ql in (it["applicability"] or "").lower()]

    stock = get_warehouse_stock(warehouse_id)

    return templates.TemplateResponse(
        "admin_stock.html",
        {
            "request": request, "user": user,
            "warehouses": warehouses, "warehouse": wh,
            "items": items, "stock": stock, "query": q,
            "all_users": list_users(),
        },
    )


@router.post("/admin/stock/income")
async def admin_stock_income(
    request: Request,
    catalog_id: int = Form(...),
    warehouse_id: int = Form(...),
    quantity: str = Form(...),
    note: str = Form(""),
):
    user = require_admin(request)
    try:
        qty = float(quantity.replace(",", "."))
        if qty <= 0:
            raise ValueError
    except ValueError:
        raise HTTPException(400, "Некорректное количество")
    add_operation(user["id"], catalog_id, warehouse_id, qty, "income", note or None)
    return RedirectResponse(f"/admin/stock?warehouse_id={warehouse_id}", status_code=303)


@router.post("/admin/stock/issue")
async def admin_stock_issue(
    request: Request,
    catalog_id: int = Form(...),
    warehouse_id: int = Form(...),
    to_user_id: int = Form(...),
    quantity: str = Form(...),
    note: str = Form(""),
):
    user = require_admin(request)
    try:
        qty = float(quantity.replace(",", "."))
        if qty <= 0:
            raise ValueError
    except ValueError:
        raise HTTPException(400, "Некорректное количество")
    add_operation(to_user_id, catalog_id, warehouse_id, qty, "take", note or None)
    return RedirectResponse(f"/admin/stock?warehouse_id={warehouse_id}", status_code=303)


@router.post("/admin/stock/return")
async def admin_stock_return(
    request: Request,
    catalog_id: int = Form(...),
    warehouse_id: int = Form(...),
    from_user_id: int = Form(...),
    quantity: str = Form(...),
    note: str = Form(""),
):
    user = require_admin(request)
    try:
        qty = float(quantity.replace(",", "."))
        if qty <= 0:
            raise ValueError
    except ValueError:
        raise HTTPException(400, "Некорректное количество")
    add_operation(from_user_id, catalog_id, warehouse_id, qty, "return", note or None)
    return RedirectResponse(f"/admin/stock?warehouse_id={warehouse_id}", status_code=303)


@router.post("/admin/stock/writeoff")
async def admin_stock_writeoff(
    request: Request,
    catalog_id: int = Form(...),
    warehouse_id: int = Form(...),
    quantity: str = Form(...),
    note: str = Form(""),
):
    user = require_admin(request)
    try:
        qty = float(quantity.replace(",", "."))
        if qty <= 0:
            raise ValueError
    except ValueError:
        raise HTTPException(400, "Некорректное количество")
    add_operation(user["id"], catalog_id, warehouse_id, qty, "writeoff", note or None)
    return RedirectResponse(f"/admin/stock?warehouse_id={warehouse_id}", status_code=303)
# =============== АДМИН: ИНВЕНТАРИЗАЦИЯ ===============

@router.get("/admin/inventory", response_class=HTMLResponse)
async def admin_inventory(request: Request, warehouse_id: int = 0):
    user = require_admin(request)
    warehouses = list_warehouses()
    if not warehouse_id and warehouses:
        warehouse_id = warehouses[0]["id"]

    wh = get_warehouse(warehouse_id) if warehouse_id else None
    if not wh:
        return templates.TemplateResponse(
            "admin_inventory.html",
            {
                "request": request, "user": user,
                "warehouses": warehouses, "warehouse": None,
                "items": [], "stock": {},
            },
        )

    items = list_catalog_for_warehouse(warehouse_id, limit=1000)
    stock = get_warehouse_stock(warehouse_id)

    return templates.TemplateResponse(
        "admin_inventory.html",
        {
            "request": request, "user": user,
            "warehouses": warehouses, "warehouse": wh,
            "items": items, "stock": stock,
        },
    )


@router.post("/admin/inventory")
async def admin_inventory_submit(request: Request, warehouse_id: int = Form(...)):
    user = require_admin(request)
    wh = get_warehouse(warehouse_id)
    if not wh:
        raise HTTPException(404, "Склад не найден")

    form = await request.form()
    stock = get_warehouse_stock(warehouse_id)

    changes = []
    for key, value in form.items():
        if not key.startswith("fact_"):
            continue
        try:
            catalog_id = int(key[5:])
        except ValueError:
            continue
        raw = (value or "").strip().replace(",", ".")
        if raw == "":
            continue
        try:
            fact = float(raw)
        except ValueError:
            continue
        current = stock.get(catalog_id, 0)
        delta = fact - current
        if abs(delta) < 1e-9:
            continue
        note = form.get("note_" + str(catalog_id), "")
        changes.append((catalog_id, delta, note or None))

    if changes:
        add_inventory_bulk(warehouse_id, changes, user["id"])

    return RedirectResponse(f"/admin/stock?warehouse_id={warehouse_id}", status_code=303)


# =============== РЮКЗАК ===============

@router.get("/backpack", response_class=HTMLResponse)
async def backpack_page(request: Request, q: str = "", tab: str = "materials"):
    user = require_user(request)

    items = get_user_backpack(user["id"])
    if q:
        ql = q.lower()
        items = [it for it in items
                 if ql in (it["name"] or "").lower()
                 or ql in (it["short_name"] or "").lower()
                 or ql in (it["article"] or "").lower()]

    tools = get_tools_by_holder(user["id"])

    incoming = list_incoming_transfers(user["id"])
    outgoing = list_outgoing_transfers(user["id"])

    return templates.TemplateResponse(
        "backpack.html",
        {
            "request": request, "user": user,
            "items": items, "query": q,
            "tools": tools,
            "tab": tab if tab in ("materials", "tools") else "materials",
            "incoming": incoming, "outgoing": outgoing,
        },
    )


@router.get("/backpack/return", response_class=HTMLResponse)
async def backpack_return_page(request: Request):
    user = require_user(request)
    user_cities = get_user_cities(user["id"])
    if not user_cities:
        raise HTTPException(400, "Вам не назначен ни один город")

    all_wh = list_warehouses()
    wh_by_city = {}
    for c in user_cities:
        wh_by_city[c["id"]] = [w for w in all_wh if w["city_id"] == c["id"]]

    items = get_user_backpack(user["id"])

    return templates.TemplateResponse(
        "backpack_return.html",
        {
            "request": request, "user": user,
            "cities": user_cities, "warehouses_by_city": wh_by_city,
            "items": items,
        },
    )


@router.post("/backpack/return")
async def backpack_return_submit(request: Request, warehouse_id: int = Form(...)):
    user = require_user(request)
    wh = get_warehouse(warehouse_id)
    if not wh:
        raise HTTPException(404, "Склад не найден")

    form = await request.form()
    cart_raw = form.get("cart", "")
    if not cart_raw:
        return RedirectResponse("/backpack", status_code=303)

    try:
        cart = json.loads(cart_raw)
    except Exception:
        raise HTTPException(400, "Некорректные данные")

    have = {b["catalog_id"]: b["balance"] for b in get_user_backpack(user["id"])}

    items = []
    for k, v in cart.items():
        try:
            cid = int(k)
            qty = float(v)
            if qty <= 0 or cid not in have:
                continue
            if qty > have[cid]:
                qty = have[cid]
            items.append((cid, qty))
        except (ValueError, TypeError):
            continue

    if items:
        add_operations_bulk(user["id"], warehouse_id, items, op_type="return")

    return RedirectResponse("/backpack?returned=1&tab=materials", status_code=303)


@router.get("/backpack/give", response_class=HTMLResponse)
async def backpack_give_page(request: Request, catalog_id: int):
    user = require_user(request)
    backpack = {b["catalog_id"]: b for b in get_user_backpack(user["id"])}
    have = backpack.get(catalog_id)
    if not have:
        raise HTTPException(400, "У вас нет этой позиции в рюкзаке")
    recipients = [u for u in list_users() if u["is_active"] and u["id"] != user["id"]]
    return templates.TemplateResponse(
        "backpack_give.html",
        {
            "request": request, "user": user,
            "item": have, "have": have,
            "recipients": recipients, "error": None,
        },
    )


@router.post("/backpack/give")
async def backpack_give_submit(
    request: Request,
    catalog_id: int = Form(...),
    to_user_id: int = Form(...),
    quantity: str = Form(...),
):
    user = require_user(request)
    backpack = {b["catalog_id"]: b for b in get_user_backpack(user["id"])}
    have = backpack.get(catalog_id)
    recipients = [u for u in list_users() if u["is_active"] and u["id"] != user["id"]]

    if not have:
        raise HTTPException(400, "У вас нет этой позиции")
    if to_user_id == user["id"]:
        raise HTTPException(400, "Нельзя передать самому себе")
    try:
        qty = float(quantity.replace(",", "."))
        if qty <= 0:
            raise ValueError
    except ValueError:
        return templates.TemplateResponse(
            "backpack_give.html",
            {
                "request": request, "user": user, "item": have, "have": have,
                "recipients": recipients, "error": "Некорректное количество",
            },
            status_code=400,
        )
    if qty > have["balance"]:
        return templates.TemplateResponse(
            "backpack_give.html",
            {
                "request": request, "user": user, "item": have, "have": have,
                "recipients": recipients,
                "error": "Доступно только " + str(have["balance"]) + " " + have["unit"],
            },
            status_code=400,
        )
    create_transfer(user["id"], to_user_id, catalog_id, qty)
    return RedirectResponse("/backpack?sent=1&tab=materials", status_code=303)


@router.post("/transfers/{transfer_id}/accept")
async def transfer_accept(request: Request, transfer_id: int):
    user = require_user(request)
    t = get_transfer(transfer_id)
    if not t or t["to_user_id"] != user["id"]:
        raise HTTPException(403, "Нет доступа")
    ok = accept_transfer(transfer_id)
    if not ok:
        raise HTTPException(400, "Не удалось принять")
    return RedirectResponse("/backpack?accepted=1&tab=materials", status_code=303)


@router.post("/transfers/{transfer_id}/reject")
async def transfer_reject(request: Request, transfer_id: int):
    user = require_user(request)
    t = get_transfer(transfer_id)
    if not t or t["to_user_id"] != user["id"]:
        raise HTTPException(403, "Нет доступа")
    reject_transfer(transfer_id)
    return RedirectResponse("/backpack?rejected=1&tab=materials", status_code=303)


@router.post("/transfers/{transfer_id}/cancel")
async def transfer_cancel(request: Request, transfer_id: int):
    user = require_user(request)
    cancel_transfer(transfer_id, user["id"])
    return RedirectResponse("/backpack?cancelled=1&tab=materials", status_code=303)


# =============== РАСХОД ИЗ РЮКЗАКА ===============

@router.get("/backpack/usage", response_class=HTMLResponse)
async def backpack_usage_page(request: Request):
    user = require_user(request)
    items = get_user_backpack(user["id"])
    return templates.TemplateResponse(
        "backpack_usage.html",
        {"request": request, "user": user, "items": items},
    )


@router.post("/backpack/usage")
async def backpack_usage_submit(request: Request):
    user = require_user(request)

    form = await request.form()
    cart_raw = form.get("cart", "")
    note = form.get("note", "") or None

    if not cart_raw:
        return RedirectResponse("/backpack?tab=materials", status_code=303)

    try:
        cart = json.loads(cart_raw)
    except Exception:
        raise HTTPException(400, "Некорректные данные")

    have = {b["catalog_id"]: b["balance"] for b in get_user_backpack(user["id"])}

    items = []
    for k, v in cart.items():
        try:
            cid = int(k)
            qty = float(v)
            if qty <= 0 or cid not in have:
                continue
            if qty > have[cid]:
                qty = have[cid]
            items.append((cid, qty))
        except (ValueError, TypeError):
            continue

    for cid, qty in items:
        item = get_catalog_item(cid)
        wh_id = item["warehouse_id"] if item else 0
        add_operation(user["id"], cid, wh_id, qty, "usage", note)

    return RedirectResponse("/backpack?used=1&tab=materials", status_code=303)


# =============== ИНСТРУМЕНТЫ ===============

@router.get("/tools", response_class=HTMLResponse)
async def tools_page(request: Request, city_id: int = 0,
                     warehouse_id: int = 0, status: str = "", q: str = ""):
    user = require_user(request)

    user_cities = get_user_cities(user["id"])
    user_city_ids = [c["id"] for c in user_cities]

    filter_city = city_id if city_id else None
    filter_wh = warehouse_id if warehouse_id else None
    filter_status = status if status in ("stock", "taken", "repair", "lost", "decommissioned") else None

    all_tools = list_tools(city_id=filter_city, warehouse_id=filter_wh, status=filter_status)

    if user["role"] != "admin":
        all_tools = [t for t in all_tools if t["city_id"] in user_city_ids]

    if q:
        ql = q.lower()
        all_tools = [t for t in all_tools
                     if ql in (t["name"] or "").lower()
                     or ql in (t["inventory_no"] or "").lower()
                     or ql in (t.get("serial_no") or "").lower()]

    return templates.TemplateResponse(
        "tools.html",
        {
            "request": request, "user": user,
            "tools": all_tools,
            "cities": list_cities() if user["role"] == "admin" else user_cities,
            "warehouses": list_warehouses(),
            "city_id": city_id, "warehouse_id": warehouse_id,
            "status": status or "",
            "count": len(all_tools),
        },
    )


@router.get("/tools/print-labels", response_class=HTMLResponse)
async def tools_print_labels(request: Request):
    user = require_admin(request)
    ids_raw = request.query_params.getlist("ids")

    cols = int(get_setting("label_cols_tool", "2") or 2)
    scale = get_setting("label_scale_tool", "medium") or "medium"

    all_tools = list_tools()
    if ids_raw:
        ids = []
        for v in ids_raw:
            try:
                ids.append(int(v))
            except ValueError:
                continue
        if ids:
            all_tools = [t for t in all_tools if t["id"] in ids]
        else:
            all_tools = []

    labels = []
    for t in all_tools:
        t = dict(t)
        t["qr_code"] = make_tool_code(t["id"])
        labels.append(t)

    return templates.TemplateResponse(
        "print_labels_tools.html",
        {
            "request": request, "user": user,
            "labels": labels,
            "label_cols": cols, "label_scale": scale,
        },
    )


@router.get("/tools/{tool_id}", response_class=HTMLResponse)
async def tool_detail(request: Request, tool_id: int):
    user = require_user(request)
    tool = get_tool(tool_id)
    if not tool:
        raise HTTPException(404, "Инструмент не найден")
    log = get_tool_log(tool_id, limit=100)
    return templates.TemplateResponse(
        "tool_detail.html",
        {"request": request, "user": user, "tool": tool, "log": log},
    )


@router.post("/tools/{tool_id}/take")
async def tool_take(request: Request, tool_id: int):
    user = require_user(request)
    ok = take_tool(tool_id, user["id"])
    if not ok:
        raise HTTPException(400, "Инструмент занят или недоступен")
    return RedirectResponse(f"/tools/{tool_id}?taken=1", status_code=303)


@router.post("/tools/{tool_id}/return")
async def tool_return(request: Request, tool_id: int):
    user = require_user(request)
    tool = get_tool(tool_id)
    if not tool:
        raise HTTPException(404, "Инструмент не найден")

    if tool["holder_id"] != user["id"] and user["role"] != "admin":
        raise HTTPException(403, "Инструмент у другого сотрудника")

    ok = return_tool(tool_id, user["id"])
    if not ok:
        raise HTTPException(400, "Не удалось вернуть")
    return RedirectResponse(f"/tools/{tool_id}?returned=1", status_code=303)


@router.post("/tools/{tool_id}/repair")
async def tool_repair(request: Request, tool_id: int):
    user = require_user(request)
    ok = repair_tool(tool_id, user["id"])
    if not ok:
        raise HTTPException(400, "Не удалось отправить в ремонт")
    return RedirectResponse(f"/tools/{tool_id}?repair=1", status_code=303)


@router.post("/tools/{tool_id}/lost")
async def tool_lost(request: Request, tool_id: int):
    user = require_admin(request)
    mark_tool_lost(tool_id, user["id"])
    return RedirectResponse(f"/tools/{tool_id}?lost=1", status_code=303)


@router.post("/tools/{tool_id}/decommission")
async def tool_decommission(request: Request, tool_id: int):
    user = require_admin(request)
    decommission_tool(tool_id, user["id"])
    return RedirectResponse(f"/tools/{tool_id}?decommissioned=1", status_code=303)


@router.get("/tools/{tool_id}/give", response_class=HTMLResponse)
async def tool_give_page(request: Request, tool_id: int):
    user = require_user(request)
    tool = get_tool(tool_id)
    if not tool:
        raise HTTPException(404, "Инструмент не найден")
    if tool["holder_id"] != user["id"]:
        raise HTTPException(403, "Инструмент не у вас")
    recipients = [u for u in list_users() if u["is_active"] and u["id"] != user["id"]]
    return templates.TemplateResponse(
        "tool_give.html",
        {
            "request": request, "user": user,
            "tool": tool, "recipients": recipients, "error": None,
        },
    )


@router.post("/tools/{tool_id}/give")
async def tool_give_submit(
    request: Request,
    tool_id: int,
    to_user_id: int = Form(...),
):
    user = require_user(request)
    ok = create_tool_transfer(tool_id, to_user_id, user["id"])
    if not ok:
        raise HTTPException(400, "Не удалось передать")
    return RedirectResponse(f"/tools/{tool_id}?given=1", status_code=303)


# =============== АДМИН: УПРАВЛЕНИЕ ИНСТРУМЕНТАМИ ===============

@router.get("/admin/tools/new", response_class=HTMLResponse)
async def tool_new_page(request: Request, city_id: int = 0):
    user = require_admin(request)
    return templates.TemplateResponse(
        "tool_form.html",
        {
            "request": request, "user": user,
            "edit_tool": None, "preselected_city": city_id,
            "cities": list_cities(),
            "warehouses": list_warehouses(),
            "next_inv": "",
            "error": None,
        },
    )


@router.post("/admin/tools/new")
async def tool_new_submit(
    request: Request,
    city_id: int = Form(...),
    warehouse_id: int = Form(...),
    inventory_no: str = Form(""),
    name: str = Form(...),
    serial_no: str = Form(""),
    note: str = Form(""),
):
    user = require_admin(request)

    if not inventory_no.strip():
        inventory_no = next_inventory_no(city_id, name)

    try:
        tool_id = create_tool(city_id, warehouse_id, inventory_no,
                              name, serial_no or None, note or None)
    except Exception as e:
        return templates.TemplateResponse(
            "tool_form.html",
            {
                "request": request, "user": user,
                "edit_tool": None, "preselected_city": city_id,
                "cities": list_cities(),
                "warehouses": list_warehouses(),
                "next_inv": inventory_no,
                "error": "Ошибка: " + str(e),
            },
            status_code=400,
        )
    return RedirectResponse(f"/tools/{tool_id}", status_code=303)


@router.get("/admin/tools/{tool_id}/edit", response_class=HTMLResponse)
async def tool_edit_page(request: Request, tool_id: int):
    user = require_admin(request)
    tool = get_tool(tool_id)
    if not tool:
        raise HTTPException(404, "Инструмент не найден")
    return templates.TemplateResponse(
        "tool_form.html",
        {
            "request": request, "user": user,
            "edit_tool": tool,
            "preselected_city": tool["city_id"],
            "cities": list_cities(),
            "warehouses": list_warehouses(),
            "next_inv": tool["inventory_no"],
            "error": None,
        },
    )


@router.post("/admin/tools/{tool_id}/edit")
async def tool_edit_submit(
    request: Request,
    tool_id: int,
    city_id: int = Form(...),
    warehouse_id: int = Form(...),
    inventory_no: str = Form(...),
    name: str = Form(...),
    serial_no: str = Form(""),
    note: str = Form(""),
):
    require_admin(request)
    update_tool(tool_id, city_id, warehouse_id, inventory_no, name,
                serial_no or None, note or None)
    return RedirectResponse(f"/tools/{tool_id}", status_code=303)


@router.post("/admin/tools/{tool_id}/delete")
async def tool_delete(request: Request, tool_id: int):
    require_admin(request)
    delete_tool(tool_id)
    return RedirectResponse("/tools", status_code=303)


@router.post("/admin/tools/{tool_id}/photo")
async def tool_photo_upload(request: Request, tool_id: int, file: UploadFile = File(...)):
    require_admin(request)
    tool = get_tool(tool_id)
    if not tool:
        raise HTTPException(404, "Инструмент не найден")

    data = await file.read()
    try:
        filename = save_tool_photo(tool_id, data)
    except Exception as e:
        raise HTTPException(400, "Ошибка обработки фото: " + str(e))

    set_tool_photo(tool_id, filename)
    return RedirectResponse(f"/tools/{tool_id}?photo=1", status_code=303)
# ============================================================
# =============== API ДЛЯ ESP32 — РЕЗОЛВ QR ==================
# ============================================================

@router.get("/api/resolve/{code}")
async def api_resolve_qr(code: str):
    entity_type, entity_id = parse_qr_code(code)

    if not entity_type or not entity_id:
        return {"ok": False, "error": "Неизвестный код"}

    if entity_type == "material":
        item = get_catalog_item(entity_id)
        if not item:
            return {"ok": False, "error": "Материал не найден"}
        return {
            "ok": True,
            "type": "material",
            "id": entity_id,
            "name": item["short_name"] or item["name"],
            "url": f"/take?warehouse_id={item['warehouse_id']}",
        }

    if entity_type == "tool":
        tool = get_tool(entity_id)
        if not tool:
            return {"ok": False, "error": "Инструмент не найден"}
        return {
            "ok": True,
            "type": "tool",
            "id": entity_id,
            "name": tool["name"],
            "url": f"/tools/{entity_id}",
        }

    if entity_type == "legacy":
        return {"ok": False, "error": "Устаревший формат этикетки. Перепечатайте."}

    return {"ok": False, "error": "Неизвестный тип"}


# =============== QR-КОДЫ ===============

@router.get("/qr/{code}.png")
async def qr_png(request: Request, code: str):
    require_user(request)
    png = make_qr_png(code.upper(), size=8)
    return Response(
        content=png,
        media_type="image/png",
        headers={"Cache-Control": "public, max-age=86400"},
    )


# =============== ПЕЧАТЬ ЭТИКЕТОК МАТЕРИАЛОВ ===============

@router.get("/catalog/print-labels", response_class=HTMLResponse)
async def catalog_print_labels(request: Request):
    user = require_admin(request)
    ids_raw = request.query_params.getlist("ids")
    wh = request.query_params.get("warehouse_id") or None

    cols_param = request.query_params.get("cols")
    rows_param = request.query_params.get("rows")

    if cols_param and cols_param in ("2", "3", "4"):
        cols = int(cols_param)
        set_setting("label_cols_material", str(cols))
    else:
        cols = int(get_setting("label_cols_material", "2") or 2)

    if rows_param and rows_param in ("4", "6", "8", "10"):
        rows = int(rows_param)
        set_setting("label_rows_material", str(rows))
    else:
        rows = int(get_setting("label_rows_material", "6") or 6)

    ids = []
    for v in ids_raw:
        try:
            ids.append(int(v))
        except ValueError:
            continue

    ids_query = "&".join("ids=" + str(x) for x in ids)
    ctx = {
        "request": request, "user": user,
        "warehouse_id": int(wh) if wh else None,
        "label_type": "material",
        "label_cols": cols, "label_rows": rows,
        "ids_query": ids_query,
    }

    if not ids:
        ctx["labels"] = []
        return templates.TemplateResponse("print_labels.html", ctx)

    items = get_catalog_items_by_ids(ids)
    labels = []
    for it in items:
        it = dict(it)
        it["qr_code"] = make_material_code(it["id"])
        labels.append(it)
    ctx["labels"] = labels

    return templates.TemplateResponse("print_labels.html", ctx)


# ============================================================
# =============== ШКАФЧИКИ С КЛЮЧАМИ ========================
# ============================================================

@router.get("/lock/{lock_id}", response_class=HTMLResponse)
async def lock_page(request: Request, lock_id: int):
    user = require_user(request)
    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")
    log = get_lock_log(lock_id, limit=50)
    cooldown_left = get_lock_cooldown_left(lock_id, cooldown_sec=30)
    seen_age = get_lock_seen_age(lock_id)
    online = seen_age is not None and seen_age < 30

    has_alarm = False
    for e in log[:20]:
        if e["event_type"] == "unauthorized_open":
            has_alarm = True
            break

    return templates.TemplateResponse(
        "lock.html",
        {
            "request": request, "user": user,
            "lock": lock, "log": log,
            "cooldown_left": cooldown_left,
            "seen_age": seen_age,
            "online": online,
            "has_alarm": has_alarm,
        },
    )


@router.post("/lock/{lock_id}/open")
async def lock_open(request: Request, lock_id: int):
    user = require_user(request)
    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")

    timings = get_lock_timings(lock_id)
    cooldown_sec = timings["cooldown_sec"]

    left = get_lock_cooldown_left(lock_id, cooldown_sec=cooldown_sec)
    if left > 0:
        return RedirectResponse(f"/lock/{lock_id}?cooldown={left}", status_code=303)

    if lock["pending_open"] or lock["is_open"]:
        return RedirectResponse(f"/lock/{lock_id}", status_code=303)

    request_lock_open(lock_id, user["id"])
    return RedirectResponse(f"/lock/{lock_id}?sent=1", status_code=303)


@router.post("/lock/{lock_id}/cancel")
async def lock_cancel(request: Request, lock_id: int):
    user = require_user(request)
    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")
    cancel_lock_open(lock_id)
    return RedirectResponse(f"/lock/{lock_id}?cancelled=1", status_code=303)


# ============ АДМИН: ШКАФЧИКИ ============

@router.get("/locks", response_class=HTMLResponse)
async def locks_list(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "locks.html",
        {"request": request, "user": user, "locks": list_locks()},
    )


@router.get("/locks/log", response_class=HTMLResponse)
async def locks_log(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "locks_log.html",
        {"request": request, "user": user, "log": get_all_lock_log(200)},
    )


@router.get("/locks/new", response_class=HTMLResponse)
async def lock_new_page(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "lock_form.html",
        {"request": request, "user": user, "edit_lock": None,
         "warehouses": list_warehouses(), "error": None},
    )


@router.post("/locks/new")
async def lock_new_submit(
    request: Request,
    warehouse_id: int = Form(...),
    name: str = Form(...),
    arduino_id: str = Form(...),
):
    user = require_admin(request)
    try:
        create_lock(warehouse_id, name, arduino_id)
    except Exception as e:
        return templates.TemplateResponse(
            "lock_form.html",
            {"request": request, "user": user, "edit_lock": None,
             "warehouses": list_warehouses(), "error": str(e)},
            status_code=400,
        )
    return RedirectResponse("/locks", status_code=303)


@router.get("/locks/{lock_id}/edit", response_class=HTMLResponse)
async def lock_edit_page(request: Request, lock_id: int):
    user = require_admin(request)
    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")
    return templates.TemplateResponse(
        "lock_form.html",
        {"request": request, "user": user, "edit_lock": lock,
         "warehouses": list_warehouses(), "error": None},
    )


@router.post("/locks/{lock_id}/edit")
async def lock_edit_submit(
    request: Request,
    lock_id: int,
    warehouse_id: int = Form(...),
    name: str = Form(...),
    arduino_id: str = Form(...),
    relay_max_ms: str = Form("7000"),
    relay_min_ms: str = Form("3000"),
    debounce_ms: str = Form("500"),
    cooldown_sec: str = Form("30"),
):
    require_admin(request)
    update_lock(lock_id, warehouse_id, name, arduino_id)

    def _int_or_none(v, min_v, max_v):
        v = (v or "").strip()
        if not v:
            return None
        try:
            n = int(v)
        except ValueError:
            return None
        return max(min_v, min(max_v, n))

    update_lock_timings(
        lock_id,
        _int_or_none(relay_max_ms, 500, 30000),
        _int_or_none(relay_min_ms, 500, 30000),
        _int_or_none(debounce_ms, 50, 5000),
        _int_or_none(cooldown_sec, 5, 3600),
    )
    return RedirectResponse("/locks", status_code=303)


@router.post("/locks/{lock_id}/delete")
async def lock_delete(request: Request, lock_id: int):
    require_admin(request)
    delete_lock(lock_id)
    return RedirectResponse("/locks", status_code=303)


# ============ API ARDUINO — ШКАФЧИК ============

@router.get("/api/arduino/poll")
async def arduino_poll(request: Request, lock_id: int = 0, door_open: int = -1):
    token = request.headers.get("X-Arduino-Token", "")
    if token != ARDUINO_TOKEN:
        raise HTTPException(403, "Неверный токен")
    if not lock_id:
        raise HTTPException(400, "lock_id обязателен")

    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")

    if door_open in (0, 1):
        if check_unauthorized_open(lock_id, door_open):
            log_lock_event(lock_id, "unauthorized_open", "Дверь открыта без команды")
            print(f"[security] Шкафчик {lock_id}: дверь открыта без команды")

    if door_open in (0, 1):
        update_lock_seen(lock_id, door_open=bool(door_open))
    else:
        update_lock_seen(lock_id)

    timings = get_lock_timings(lock_id)
    config = {
        "relay_max_ms": timings["relay_max_ms"],
        "relay_min_ms": timings["relay_min_ms"],
        "debounce_ms":  timings["debounce_ms"],
    }

    if lock["pending_open"]:
        consume_lock_open(lock_id)
        log_lock_event(lock_id, "command_sent", "Шкафчик сейчас откроется")
        return {"command": "open", "lock_id": lock_id, "config": config}

    if is_calibration_requested(lock_id):
        return {"command": "calibrate", "lock_id": lock_id, "config": config}

    return {"command": "none", "config": config}


@router.post("/api/arduino/event")
async def arduino_event(request: Request):
    token = request.headers.get("X-Arduino-Token", "")
    if token != ARDUINO_TOKEN:
        raise HTTPException(403, "Неверный токен")

    try:
        data = await request.json()
    except Exception:
        raise HTTPException(400, "Некорректный JSON")

    lock_id = data.get("lock_id")
    event = data.get("event")
    note = data.get("note", "")

    if not lock_id or event not in ("opened", "closed", "timeout"):
        raise HTTPException(400, "Нужны lock_id и event")

    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")

    user_id = lock.get("requested_by")
    log_lock_event(lock_id, event, note, user_id)
    return {"ok": True}


# ============ API ARDUINO — КЛЮЧИ ============

@router.post("/api/arduino/key_event")
async def arduino_key_event(request: Request):
    token = request.headers.get("X-Arduino-Token", "")
    if token != ARDUINO_TOKEN:
        raise HTTPException(403, "Неверный токен")

    try:
        data = await request.json()
    except Exception:
        raise HTTPException(400, "Некорректный JSON")

    lock_id = data.get("lock_id")
    slot_code = data.get("slot_code")
    is_present = data.get("is_present")
    adc = data.get("adc")

    if not lock_id or not slot_code or is_present not in (0, 1):
        raise HTTPException(400, "Нужны lock_id, slot_code, is_present")

    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")

    changed = update_slot_state(lock_id, slot_code, bool(is_present), adc)

    if changed:
        requested_by = lock.get("requested_by")
        if requested_by:
            link_key_event_to_user(lock_id, requested_by)

    return {"ok": True, "changed": changed}


@router.post("/api/arduino/key_calibration")
async def arduino_key_calibration(request: Request):
    token = request.headers.get("X-Arduino-Token", "")
    if token != ARDUINO_TOKEN:
        raise HTTPException(403, "Неверный токен")

    try:
        data = await request.json()
    except Exception:
        raise HTTPException(400, "Некорректный JSON")

    lock_id = data.get("lock_id")
    slots = data.get("slots", {})

    if not lock_id or not isinstance(slots, dict):
        raise HTTPException(400, "Нужны lock_id и slots")

    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")

    init_empty_slots(lock_id)

    updated = 0
    for slot_code, info in slots.items():
        adc = info.get("adc")
        if adc is None:
            continue
        if adc <= 5 or adc >= 4090:
            continue
        update_key_calibration(lock_id, slot_code, int(adc), None)
        updated += 1

    for slot_code, info in slots.items():
        is_present = info.get("is_present")
        adc = info.get("adc")
        if is_present is not None:
            update_slot_state(lock_id, slot_code, bool(is_present), adc)

    consume_calibration(lock_id)
    print(f"[calibration] Шкафчик {lock_id}: откалибровано {updated} гнёзд")

    return {"ok": True, "updated": updated}


# ============ АДМИН: УПРАВЛЕНИЕ КЛЮЧАМИ ============

@router.get("/admin/locks/{lock_id}/keys", response_class=HTMLResponse)
async def lock_keys_page(request: Request, lock_id: int):
    user = require_admin(request)
    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")

    init_empty_slots(lock_id)

    keys = list_keys(lock_id)
    progress = get_calibration_progress(lock_id)
    conflicts = find_calibration_conflicts(lock_id)
    stats = count_keys_in_lock(lock_id)

    rows_data = {"A": [], "B": [], "C": []}
    for k in keys:
        row_letter = k["slot_code"][0]
        if row_letter in rows_data:
            rows_data[row_letter].append(k)

    return templates.TemplateResponse(
        "lock_keys.html",
        {
            "request": request, "user": user,
            "lock": lock,
            "rows_data": rows_data,
            "progress": progress,
            "conflicts": conflicts,
            "stats": stats,
            "key_log": get_key_log(lock_id, limit=50),
        },
    )


@router.get("/admin/locks/{lock_id}/keys/{slot}/edit", response_class=HTMLResponse)
async def key_edit_page(request: Request, lock_id: int, slot: str):
    user = require_admin(request)
    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")

    key = get_key(lock_id, slot)
    if not key:
        upsert_key(lock_id, slot, None, None)
        key = get_key(lock_id, slot)

    return templates.TemplateResponse(
        "key_form.html",
        {
            "request": request, "user": user,
            "lock": lock, "key": key,
            "error": None,
        },
    )


@router.post("/admin/locks/{lock_id}/keys/{slot}/edit")
async def key_edit_submit(
    request: Request,
    lock_id: int,
    slot: str,
    name: str = Form(""),
    description: str = Form(""),
):
    require_admin(request)
    upsert_key(lock_id, slot, name.strip() or None, description.strip() or None)
    return RedirectResponse(f"/admin/locks/{lock_id}/keys", status_code=303)


@router.post("/admin/locks/{lock_id}/keys/{slot}/clear")
async def key_clear(request: Request, lock_id: int, slot: str):
    require_admin(request)
    delete_key_data(lock_id, slot)
    return RedirectResponse(f"/admin/locks/{lock_id}/keys", status_code=303)


# ============ АДМИН: КАЛИБРОВКА ============

@router.get("/admin/locks/{lock_id}/calibrate", response_class=HTMLResponse)
async def lock_calibrate_page(request: Request, lock_id: int):
    user = require_admin(request)
    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")

    seen_age = get_lock_seen_age(lock_id)
    online = seen_age is not None and seen_age < 30

    progress = get_calibration_progress(lock_id)
    conflicts = find_calibration_conflicts(lock_id)

    return templates.TemplateResponse(
        "lock_calibration.html",
        {
            "request": request, "user": user,
            "lock": lock,
            "online": online,
            "seen_age": seen_age,
            "progress": progress,
            "conflicts": conflicts,
            "started": request.query_params.get("started") == "1",
        },
    )


@router.post("/admin/locks/{lock_id}/calibrate")
async def lock_calibrate_start(request: Request, lock_id: int):
    require_admin(request)
    lock = get_lock(lock_id)
    if not lock:
        raise HTTPException(404, "Шкафчик не найден")

    keys = list_keys(lock_id)
    for k in keys:
        update_key_calibration(lock_id, k["slot_code"], None, None)

    request_calibration(lock_id)
    return RedirectResponse(f"/admin/locks/{lock_id}/calibrate?started=1", status_code=303)


@router.get("/admin/locks/{lock_id}/calibrate/status")
async def lock_calibrate_status(request: Request, lock_id: int):
    require_admin(request)
    progress = get_calibration_progress(lock_id)
    conflicts = find_calibration_conflicts(lock_id)
    return {
        "calibrated": progress["calibrated"],
        "total": progress["total"],
        "requested": progress["requested"],
        "scan_complete": progress["scan_complete"],
        "conflicts": conflicts,
    }


# =============== ОТЧЁТЫ ===============

@router.get("/reports", response_class=HTMLResponse)
async def reports_index(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse("reports.html", {"request": request, "user": user})


@router.get("/reports/stock", response_class=HTMLResponse)
async def report_stock(request: Request):
    user = require_admin(request)
    wh_id = request.query_params.get("warehouse_id") or None
    city_id = request.query_params.get("city_id") or None
    wh_id = int(wh_id) if wh_id else None
    city_id = int(city_id) if city_id else None
    rows = get_stock_report(warehouse_id=wh_id, city_id=city_id)
    return templates.TemplateResponse(
        "report_stock.html",
        {
            "request": request, "user": user, "rows": rows,
            "warehouses": list_warehouses(), "cities": list_cities(),
            "warehouse_id": wh_id, "city_id": city_id,
        },
    )


@router.get("/reports/stock/export")
async def report_stock_export(request: Request):
    require_admin(request)
    wh_id = request.query_params.get("warehouse_id") or None
    city_id = request.query_params.get("city_id") or None
    wh_id = int(wh_id) if wh_id else None
    city_id = int(city_id) if city_id else None
    rows = get_stock_report(warehouse_id=wh_id, city_id=city_id)
    return _excel_response(build_stock_excel(rows), "stock")


@router.get("/reports/backpacks", response_class=HTMLResponse)
async def report_backpacks(request: Request):
    user = require_admin(request)
    wh_id = request.query_params.get("warehouse_id") or None
    user_id = request.query_params.get("user_id") or None
    wh_id = int(wh_id) if wh_id else None
    user_id = int(user_id) if user_id else None
    rows = get_backpacks_report(warehouse_id=wh_id, user_id=user_id)
    return templates.TemplateResponse(
        "report_backpacks.html",
        {
            "request": request, "user": user, "rows": rows,
            "warehouses": list_warehouses(), "users": list_users(),
            "warehouse_id": wh_id, "user_id": user_id,
        },
    )


@router.get("/reports/backpacks/export")
async def report_backpacks_export(request: Request):
    require_admin(request)
    wh_id = request.query_params.get("warehouse_id") or None
    user_id = request.query_params.get("user_id") or None
    wh_id = int(wh_id) if wh_id else None
    user_id = int(user_id) if user_id else None
    rows = get_backpacks_report(warehouse_id=wh_id, user_id=user_id)
    return _excel_response(build_backpacks_excel(rows), "backpacks")


@router.get("/reports/operations", response_class=HTMLResponse)
async def report_operations(request: Request):
    user = require_admin(request)
    df, dt = _parse_dates(request)
    wh_id = request.query_params.get("warehouse_id") or None
    user_id = request.query_params.get("user_id") or None
    op_type = request.query_params.get("op_type") or None
    wh_id = int(wh_id) if wh_id else None
    user_id = int(user_id) if user_id else None
    rows = get_operations_report(date_from=df, date_to=dt,
                                 warehouse_id=wh_id, user_id=user_id,
                                 op_type=op_type, limit=1000)
    return templates.TemplateResponse(
        "report_operations.html",
        {
            "request": request, "user": user, "rows": rows,
            "warehouses": list_warehouses(), "users": list_users(),
            "date_from": df or "", "date_to": dt or "",
            "warehouse_id": wh_id, "user_id": user_id,
            "op_type": op_type or "",
        },
    )


@router.get("/reports/operations/export")
async def report_operations_export(request: Request):
    require_admin(request)
    df, dt = _parse_dates(request)
    wh_id = request.query_params.get("warehouse_id") or None
    user_id = request.query_params.get("user_id") or None
    op_type = request.query_params.get("op_type") or None
    wh_id = int(wh_id) if wh_id else None
    user_id = int(user_id) if user_id else None
    rows = get_operations_report(date_from=df, date_to=dt,
                                 warehouse_id=wh_id, user_id=user_id,
                                 op_type=op_type, limit=100000)
    return _excel_response(build_operations_excel(rows), "operations")


@router.get("/reports/summary-users", response_class=HTMLResponse)
async def report_summary_users(request: Request):
    user = require_admin(request)
    df, dt = _parse_dates(request)
    wh_id = request.query_params.get("warehouse_id") or None
    user_id = request.query_params.get("user_id") or None
    wh_id = int(wh_id) if wh_id else None
    user_id = int(user_id) if user_id else None
    rows = get_summary_users_report(date_from=df, date_to=dt,
                                    warehouse_id=wh_id, user_id=user_id)
    return templates.TemplateResponse(
        "report_summary_users.html",
        {
            "request": request, "user": user, "rows": rows,
            "warehouses": list_warehouses(), "users": list_users(),
            "date_from": df or "", "date_to": dt or "",
            "warehouse_id": wh_id, "user_id": user_id,
        },
    )


@router.get("/reports/summary-users/export")
async def report_summary_users_export(request: Request):
    require_admin(request)
    df, dt = _parse_dates(request)
    wh_id = request.query_params.get("warehouse_id") or None
    user_id = request.query_params.get("user_id") or None
    wh_id = int(wh_id) if wh_id else None
    user_id = int(user_id) if user_id else None
    rows = get_summary_users_report(date_from=df, date_to=dt,
                                    warehouse_id=wh_id, user_id=user_id)
    return _excel_response(build_summary_users_excel(rows), "summary_users")


@router.get("/reports/summary-items", response_class=HTMLResponse)
async def report_summary_items(request: Request):
    user = require_admin(request)
    df, dt = _parse_dates(request)
    wh_id = request.query_params.get("warehouse_id") or None
    city_id = request.query_params.get("city_id") or None
    wh_id = int(wh_id) if wh_id else None
    city_id = int(city_id) if city_id else None
    rows = get_summary_items_report(date_from=df, date_to=dt,
                                    warehouse_id=wh_id, city_id=city_id)
    return templates.TemplateResponse(
        "report_summary_items.html",
        {
            "request": request, "user": user, "rows": rows,
            "warehouses": list_warehouses(), "cities": list_cities(),
            "date_from": df or "", "date_to": dt or "",
            "warehouse_id": wh_id, "city_id": city_id,
        },
    )


@router.get("/reports/summary-items/export")
async def report_summary_items_export(request: Request):
    require_admin(request)
    df, dt = _parse_dates(request)
    wh_id = request.query_params.get("warehouse_id") or None
    city_id = request.query_params.get("city_id") or None
    wh_id = int(wh_id) if wh_id else None
    city_id = int(city_id) if city_id else None
    rows = get_summary_items_report(date_from=df, date_to=dt,
                                    warehouse_id=wh_id, city_id=city_id)
    return _excel_response(build_summary_items_excel(rows), "summary_items")


@router.get("/reports/transfers", response_class=HTMLResponse)
async def report_transfers(request: Request):
    user = require_admin(request)
    df, dt = _parse_dates(request)
    status = request.query_params.get("status") or None
    rows = get_transfers_report(date_from=df, date_to=dt, status=status)
    return templates.TemplateResponse(
        "report_transfers.html",
        {
            "request": request, "user": user, "rows": rows,
            "date_from": df or "", "date_to": dt or "", "status": status or "",
        },
    )


@router.get("/reports/transfers/export")
async def report_transfers_export(request: Request):
    require_admin(request)
    df, dt = _parse_dates(request)
    status = request.query_params.get("status") or None
    rows = get_transfers_report(date_from=df, date_to=dt, status=status)
    return _excel_response(build_transfers_excel(rows), "transfers")


def _excel_response(data: bytes, prefix: str):
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="' + export_filename(prefix) + '"'},
    )


# =============== ИСТОРИЯ ===============

@router.get("/history", response_class=HTMLResponse)
async def history(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "history.html",
        {"request": request, "user": user, "operations": list_operations(200)},
    )


@router.get("/history/{op_id}/edit", response_class=HTMLResponse)
async def operation_edit_page(request: Request, op_id: int):
    user = require_admin(request)
    op = get_operation(op_id)
    if not op:
        raise HTTPException(404, "Операция не найдена")
    return templates.TemplateResponse(
        "operation_form.html",
        {"request": request, "user": user, "op": op, "error": None},
    )


@router.post("/history/{op_id}/edit")
async def operation_edit_submit(
    request: Request,
    op_id: int,
    catalog_id: int = Form(...),
    warehouse_id: int = Form(...),
    quantity: str = Form(...),
    operation_type: str = Form(...),
    note: str = Form(""),
):
    require_admin(request)
    if operation_type not in ("take", "return", "income", "usage",
                              "writeoff", "inventory",
                              "transfer_out", "transfer_in"):
        raise HTTPException(400, "Некорректный тип")
    try:
        qty = float(quantity.replace(",", "."))
    except ValueError:
        raise HTTPException(400, "Некорректное количество")
    update_operation(op_id, catalog_id, warehouse_id, qty, operation_type, note or None)
    return RedirectResponse("/history", status_code=303)


@router.post("/history/{op_id}/delete")
async def operation_delete(request: Request, op_id: int):
    require_admin(request)
    delete_operation(op_id)
    return RedirectResponse("/history", status_code=303)


# =============== ГОРОДА ===============

@router.get("/cities", response_class=HTMLResponse)
async def cities_page(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "cities.html",
        {"request": request, "user": user, "cities": list_cities()},
    )


@router.get("/cities/new", response_class=HTMLResponse)
async def city_new_page(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "city_form.html",
        {"request": request, "user": user, "edit_city": None, "error": None},
    )


@router.post("/cities/new")
async def city_new_submit(request: Request, name: str = Form(...)):
    user = require_admin(request)
    try:
        create_city(name)
    except Exception as e:
        return templates.TemplateResponse(
            "city_form.html",
            {"request": request, "user": user, "edit_city": None,
             "error": "Ошибка: " + str(e)},
            status_code=400,
        )
    return RedirectResponse("/cities", status_code=303)


@router.get("/cities/{city_id}/edit", response_class=HTMLResponse)
async def city_edit_page(request: Request, city_id: int):
    user = require_admin(request)
    c = get_city(city_id)
    if not c:
        raise HTTPException(404, "Город не найден")
    return templates.TemplateResponse(
        "city_form.html",
        {"request": request, "user": user, "edit_city": c, "error": None},
    )


@router.post("/cities/{city_id}/edit")
async def city_edit_submit(request: Request, city_id: int, name: str = Form(...)):
    require_admin(request)
    update_city(city_id, name)
    return RedirectResponse("/cities", status_code=303)


@router.post("/cities/{city_id}/delete")
async def city_delete(request: Request, city_id: int):
    require_admin(request)
    delete_city(city_id)
    return RedirectResponse("/cities", status_code=303)


# =============== СКЛАДЫ ===============

@router.get("/warehouses", response_class=HTMLResponse)
async def warehouses_page(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "warehouses.html",
        {"request": request, "user": user, "warehouses": list_warehouses()},
    )


@router.get("/warehouses/new", response_class=HTMLResponse)
async def warehouse_new_page(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "warehouse_form.html",
        {"request": request, "user": user, "edit_wh": None,
         "cities": list_cities(), "error": None},
    )


@router.post("/warehouses/new")
async def warehouse_new_submit(
    request: Request,
    city_id: int = Form(...),
    name: str = Form(...),
    address: str = Form(""),
    room: str = Form(""),
    responsible: str = Form(""),
):
    user = require_admin(request)
    create_warehouse(city_id, name, address, room, responsible)
    return RedirectResponse("/warehouses", status_code=303)


@router.get("/warehouses/{wh_id}/edit", response_class=HTMLResponse)
async def warehouse_edit_page(request: Request, wh_id: int):
    user = require_admin(request)
    wh = get_warehouse(wh_id)
    if not wh:
        raise HTTPException(404, "Склад не найден")
    return templates.TemplateResponse(
        "warehouse_form.html",
        {"request": request, "user": user, "edit_wh": wh,
         "cities": list_cities(), "error": None},
    )


@router.post("/warehouses/{wh_id}/edit")
async def warehouse_edit_submit(
    request: Request,
    wh_id: int,
    city_id: int = Form(...),
    name: str = Form(...),
    address: str = Form(""),
    room: str = Form(""),
    responsible: str = Form(""),
):
    require_admin(request)
    update_warehouse(wh_id, city_id, name, address, room, responsible)
    return RedirectResponse("/warehouses", status_code=303)


@router.post("/warehouses/{wh_id}/delete")
async def warehouse_delete(request: Request, wh_id: int):
    require_admin(request)
    delete_warehouse(wh_id)
    return RedirectResponse("/warehouses", status_code=303)


# =============== КАТАЛОГ ===============

@router.get("/catalog", response_class=HTMLResponse)
async def catalog_page(request: Request, warehouse_id: int = 0):
    user = require_admin(request)
    wh_id = warehouse_id or None
    return templates.TemplateResponse(
        "catalog.html",
        {
            "request": request, "user": user,
            "items": list_catalog(wh_id, limit=1000),
            "warehouses": list_warehouses(),
            "selected_warehouse": warehouse_id,
            "count": catalog_count(),
        },
    )


@router.get("/catalog/new", response_class=HTMLResponse)
async def catalog_new_page(request: Request, warehouse_id: int = 0):
    user = require_admin(request)
    return templates.TemplateResponse(
        "catalog_form.html",
        {
            "request": request, "user": user,
            "edit_item": None, "warehouses": list_warehouses(),
            "locations": list_locations(warehouse_id or None),
            "preselected_warehouse": warehouse_id,
            "error": None,
        },
    )


@router.post("/catalog/new")
async def catalog_new_submit(
    request: Request,
    warehouse_id: int = Form(...),
    article: str = Form(""),
    name: str = Form(...),
    short_name: str = Form(""),
    segment: str = Form(""),
    applicability: str = Form(""),
    unit: str = Form("шт"),
    location_id: str = Form(""),
    sort_order: str = Form("1000"),
):
    user = require_admin(request)
    loc_id = int(location_id) if location_id.strip() else None
    try:
        so = int(sort_order)
    except ValueError:
        so = 1000
    try:
        create_catalog_item(
            warehouse_id, article, name, short_name or None,
            segment or None, applicability or None, unit or "шт", loc_id, so,
        )
    except Exception as e:
        return templates.TemplateResponse(
            "catalog_form.html",
            {
                "request": request, "user": user,
                "edit_item": None, "warehouses": list_warehouses(),
                "locations": list_locations(warehouse_id),
                "preselected_warehouse": warehouse_id,
                "error": "Ошибка: " + str(e),
            },
            status_code=400,
        )
    return RedirectResponse("/catalog?warehouse_id=" + str(warehouse_id), status_code=303)


@router.get("/catalog/{item_id}/edit", response_class=HTMLResponse)
async def catalog_edit_page(request: Request, item_id: int):
    user = require_admin(request)
    item = get_catalog_item(item_id)
    if not item:
        raise HTTPException(404, "Позиция не найдена")
    return templates.TemplateResponse(
        "catalog_form.html",
        {
            "request": request, "user": user,
            "edit_item": item, "warehouses": list_warehouses(),
            "locations": list_locations(item["warehouse_id"]),
            "preselected_warehouse": item["warehouse_id"],
            "error": None,
        },
    )


@router.post("/catalog/{item_id}/edit")
async def catalog_edit_submit(
    request: Request,
    item_id: int,
    warehouse_id: int = Form(...),
    article: str = Form(""),
    name: str = Form(...),
    short_name: str = Form(""),
    segment: str = Form(""),
    applicability: str = Form(""),
    unit: str = Form("шт"),
    location_id: str = Form(""),
    sort_order: str = Form("1000"),
):
    require_admin(request)
    loc_id = int(location_id) if location_id.strip() else None
    try:
        so = int(sort_order)
    except ValueError:
        so = 1000
    update_catalog_item(
        item_id, warehouse_id, article, name, short_name or None,
        segment or None, applicability or None, unit or "шт", loc_id, so,
    )
    return RedirectResponse("/catalog?warehouse_id=" + str(warehouse_id), status_code=303)


@router.post("/catalog/{item_id}/sort")
async def catalog_sort_submit(
    request: Request,
    item_id: int,
    sort_order: str = Form(...),
    warehouse_id: int = Form(...),
):
    require_admin(request)
    try:
        so = int(sort_order)
    except ValueError:
        so = 1000
    update_catalog_sort(item_id, so)
    return RedirectResponse("/catalog?warehouse_id=" + str(warehouse_id), status_code=303)


@router.post("/catalog/{item_id}/delete")
async def catalog_delete(request: Request, item_id: int):
    require_admin(request)
    item = get_catalog_item(item_id)
    wid = item["warehouse_id"] if item else 0
    delete_catalog_item(item_id)
    return RedirectResponse("/catalog?warehouse_id=" + str(wid), status_code=303)


@router.get("/catalog/template")
async def catalog_template(request: Request):
    require_admin(request)
    wb = Workbook()
    ws = wb.active
    ws.title = "Каталог"
    headers = ["Ном. Позиции", "Наименование", "Короткое имя",
               "Сегмент", "Применяемость", "Ед. изм", "Место"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.append(["016.1000.1020",
               "Болт анкерный распорный, с крюком М8 8х60мм сталь оцинкованный",
               "Болт анкерный М8", "B2B", "Инсталляции", "шт", "A-12"])
    widths = [18, 50, 22, 12, 18, 8, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="catalog_template.xlsx"'},
    )


@router.get("/catalog/import", response_class=HTMLResponse)
async def catalog_import_page(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "catalog_import.html",
        {
            "request": request, "user": user,
            "warehouses": list_warehouses(), "result": None, "error": None,
        },
    )


@router.post("/catalog/import")
async def catalog_import_submit(
    request: Request,
    warehouse_id: int = Form(...),
    file: UploadFile = File(...),
):
    import sqlite3
    user = require_admin(request)
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        return templates.TemplateResponse(
            "catalog_import.html",
            {
                "request": request, "user": user,
                "warehouses": list_warehouses(),
                "result": None, "error": "Нужен файл .xlsx или .xlsm",
            },
            status_code=400,
        )
    try:
        wb = load_workbook(file.file, data_only=True)
        ws = wb.active
        alias = {
            "article":       ["Ном. Позиции", "Ном позиции", "Артикул"],
            "name":          ["Наименование", "Наименование позиции"],
            "short_name":    ["Короткое имя", "Короткое наименование", "Кратко"],
            "segment":       ["Сегмент"],
            "applicability": ["Применяемость"],
            "unit":          ["Ед. изм", "Ед.изм", "Единица"],
            "location":      ["Место", "Локация", "Ячейка"],
        }
        header_row = None
        cols = {}
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
            if not row:
                continue
            cells = [str(c).strip() if c is not None else "" for c in row]
            has_art = any(any(a.lower() == c.lower() for a in alias["article"]) for c in cells)
            has_name = any(any(a.lower() == c.lower() for a in alias["name"]) for c in cells)
            if has_art and has_name:
                header_row = i
                for field, aliases in alias.items():
                    for idx, cell in enumerate(cells):
                        if any(a.lower() == cell.lower() for a in aliases):
                            cols[field] = idx
                            break
                break
        if not header_row:
            raise ValueError("Не нашёл строку с заголовками")

        loc_cache = {l["name"]: l["id"] for l in list_locations(warehouse_id)}

        def get_loc_id(name):
            if not name:
                return None
            name = name.strip()
            if not name:
                return None
            if name in loc_cache:
                return loc_cache[name]
            create_location(warehouse_id, name)
            new_loc = next(
                (l for l in list_locations(warehouse_id) if l["name"] == name), None
            )
            if new_loc:
                loc_cache[name] = new_loc["id"]
                return new_loc["id"]
            return None

        inserted, skipped = 0, 0
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row:
                continue

            def val(field):
                idx = cols.get(field)
                if idx is None or idx >= len(row) or row[idx] is None:
                    return ""
                return str(row[idx]).strip()

            article = val("article")
            name = val("name")
            if not name:
                continue

            try:
                create_catalog_item(
                    warehouse_id=warehouse_id,
                    article=article,
                    name=name,
                    short_name=val("short_name") or None,
                    segment=val("segment") or None,
                    applicability=val("applicability") or None,
                    unit=val("unit") or "шт",
                    location_id=get_loc_id(val("location")),
                    sort_order=1000,
                )
                inserted += 1
            except sqlite3.IntegrityError:
                skipped += 1
            except Exception as e:
                print(f"[import] ERROR {name}: {type(e).__name__}: {e}")
                skipped += 1

        return templates.TemplateResponse(
            "catalog_import.html",
            {
                "request": request, "user": user,
                "warehouses": list_warehouses(),
                "result": {"inserted": inserted, "skipped": skipped,
                           "total": catalog_count()},
                "error": None,
            },
        )
    except Exception as e:
        return templates.TemplateResponse(
            "catalog_import.html",
            {
                "request": request, "user": user,
                "warehouses": list_warehouses(),
                "result": None, "error": "Ошибка: " + str(e),
            },
            status_code=400,
        )


# =============== НАСТРОЙКИ ===============

@router.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "settings.html",
        {
            "request": request, "user": user,
            "settings": get_all_settings(),
        },
    )


@router.post("/settings")
async def settings_submit(request: Request):
    require_admin(request)
    form = await request.form()
    set_setting("show_favorites", "1" if form.get("show_favorites") else "0")
    set_setting("show_recent", "1" if form.get("show_recent") else "0")
    set_setting("show_stock", "1" if form.get("show_stock") else "0")
    set_setting("recent_limit", str(_safe_int(form.get("recent_limit"), 10, 1, 50)))

    set_setting("lock_cooldown_sec", str(_safe_int(form.get("lock_cooldown_sec"), 30, 5, 3600)))
    set_setting("lock_relay_pulse_max_ms", str(_safe_int(form.get("lock_relay_pulse_max_ms"), 7000, 500, 30000)))
    set_setting("lock_relay_pulse_min_ms", str(_safe_int(form.get("lock_relay_pulse_min_ms"), 3000, 500, 30000)))
    set_setting("lock_debounce_ms", str(_safe_int(form.get("lock_debounce_ms"), 500, 50, 5000)))
    set_setting("lock_open_timeout_sec", str(_safe_int(form.get("lock_open_timeout_sec"), 30, 5, 600)))
    set_setting("tool_hold_days_alert", str(_safe_int(form.get("tool_hold_days_alert"), 0, 0, 365)))

    return RedirectResponse("/settings", status_code=303)


# =============== ПОЛЬЗОВАТЕЛИ ===============

@router.get("/users", response_class=HTMLResponse)
async def users_page(request: Request):
    user = require_admin(request)
    users = list_users()
    for u in users:
        u["cities"] = get_user_cities(u["id"])
    return templates.TemplateResponse(
        "users.html",
        {"request": request, "user": user, "users": users},
    )


@router.get("/users/new", response_class=HTMLResponse)
async def user_new_page(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "user_form.html",
        {
            "request": request, "user": user,
            "edit_user": None, "cities": list_cities(),
            "user_city_ids": [], "error": None,
        },
    )


@router.post("/users/new")
async def user_new_submit(
    request: Request,
    full_name: str = Form(...),
    username: str = Form(...),
    password: str = Form(...),
    role: str = Form("user"),
):
    user = require_admin(request)
    if role not in ("user", "admin"):
        role = "user"
    form = await request.form()
    city_ids = [int(x) for x in form.getlist("city_ids") if x]
    try:
        create_user(full_name.strip(), username.strip(),
                    hash_password(password), role, city_ids)
    except Exception as e:
        return templates.TemplateResponse(
            "user_form.html",
            {
                "request": request, "user": user, "edit_user": None,
                "cities": list_cities(), "user_city_ids": city_ids,
                "error": "Ошибка: " + str(e),
            },
            status_code=400,
        )
    return RedirectResponse("/users", status_code=303)


@router.get("/users/{user_id}/edit", response_class=HTMLResponse)
async def user_edit_page(request: Request, user_id: int):
    user = require_admin(request)
    target = get_user_by_id(user_id)
    if not target:
        raise HTTPException(404, "Пользователь не найден")
    target["cities"] = get_user_cities(user_id)
    return templates.TemplateResponse(
        "user_form.html",
        {
            "request": request, "user": user,
            "edit_user": target, "cities": list_cities(),
            "user_city_ids": [c["id"] for c in target["cities"]],
            "error": None,
        },
    )


@router.post("/users/{user_id}/edit")
async def user_edit_submit(
    request: Request,
    user_id: int,
    full_name: str = Form(...),
    role: str = Form("user"),
    is_active: str = Form(""),
    password: str = Form(""),
):
    require_admin(request)
    if role not in ("user", "admin"):
        role = "user"
    form = await request.form()
    city_ids = [int(x) for x in form.getlist("city_ids") if x]
    pwd_hash = hash_password(password) if password.strip() else None
    update_user(
        user_id,
        full_name=full_name.strip(),
        role=role,
        is_active=(is_active == "on"),
        password_hash=pwd_hash,
        city_ids=city_ids,
    )
    return RedirectResponse("/users", status_code=303)


@router.post("/users/{user_id}/delete")
async def user_delete(request: Request, user_id: int):
    user = require_admin(request)
    if user["id"] == user_id:
        raise HTTPException(400, "Нельзя удалить самого себя")
    delete_user(user_id)
    return RedirectResponse("/users", status_code=303)


# =============== СБРОС ===============

@router.get("/reset", response_class=HTMLResponse)
async def reset_page(request: Request):
    user = require_admin(request)
    return templates.TemplateResponse(
        "reset.html",
        {"request": request, "user": user, "error": None},
    )


@router.post("/reset")
async def reset_operations(
    request: Request,
    confirm1: str = Form(""),
    confirm2: str = Form(""),
    confirm3: str = Form(""),
):
    user = require_admin(request)
    if not (confirm1 == "on" and confirm2 == "on" and confirm3 == "on"):
        return templates.TemplateResponse(
            "reset.html",
            {
                "request": request, "user": user,
                "error": "Нужно включить все три подтверждения.",
            },
            status_code=400,
        )
    clear_operations()
    return RedirectResponse("/", status_code=303)


# =============== РЕЗЕРВНЫЕ КОПИИ ===============

from core.backup import (
    create_backup, list_backups, auto_backup_if_needed,
    BACKUP_DIR,
)


@router.get("/admin/backups", response_class=HTMLResponse)
async def backups_page(request: Request):
    user = require_admin(request)
    backups = list_backups()
    return templates.TemplateResponse(
        "backups.html",
        {
            "request": request, "user": user,
            "backups": backups,
            "keep_count": 7,
        },
    )


@router.post("/admin/backups/now")
async def backup_now(request: Request):
    require_admin(request)
    create_backup(tag="manual")
    return RedirectResponse("/admin/backups?done=1", status_code=303)


@router.get("/admin/backups/{filename}")
async def backup_download(request: Request, filename: str):
    require_admin(request)
    if "/" in filename or ".." in filename or not filename.startswith("warehouse_"):
        raise HTTPException(400, "Некорректное имя")
    path = os.path.join(BACKUP_DIR, filename)
    if not os.path.exists(path):
        raise HTTPException(404, "Файл не найден")
    return FileResponse(
        path,
        media_type="application/octet-stream",
        filename=filename,
    )


@router.post("/admin/backups/{filename}/delete")
async def backup_delete(request: Request, filename: str):
    require_admin(request)
    if "/" in filename or ".." in filename or not filename.startswith("warehouse_"):
        raise HTTPException(400, "Некорректное имя")
    path = os.path.join(BACKUP_DIR, filename)
    if os.path.exists(path):
        os.remove(path)
    return RedirectResponse("/admin/backups", status_code=303)